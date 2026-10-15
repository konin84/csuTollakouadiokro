"""Exports Excel des rapports (mise en forme proche de l'ancien tableau)."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="1F4E46")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="E7EFEC")
THIN = Side(style="thin", color="C9D3D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM = '#,##0;[Red]-#,##0;""'
NUM0 = '#,##0;[Red]-#,##0;0'


def _title(ws, title, subtitle):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, color="555555")


def _table(ws, row, headers, rows, total=None, widths=None, num_from=1, zero=False):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row, c, h)
        cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[row].height = 32
    r = row
    for values in rows:
        r += 1
        for c, v in enumerate(values, start=1):
            cell = ws.cell(r, c, v)
            cell.border = BORDER
            if c > num_from and isinstance(v, (int, float)):
                cell.number_format = NUM0 if zero else NUM
    if total:
        r += 1
        for c, v in enumerate(total, start=1):
            cell = ws.cell(r, c, v)
            cell.font, cell.fill, cell.border = Font(bold=True), TOTAL_FILL, BORDER
            if isinstance(v, (int, float)):
                cell.number_format = NUM0
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row + 1, 2)
    return r


def _save(wb):
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def monthly_workbook(d):
    wb = Workbook()
    sub = f"{d['facility'].name} – {d['month_name']}"
    t = d["total"]

    ws = wb.active
    ws.title = "Synthèse journalière"
    _title(ws, "TABLEAU DE BORD RECETTES", sub)
    headers = ["Date", "Entrées (hors labo)", "Laboratoire", "Échographie", "Pharmacie", "Recette brute",
               "Commissions labo reversées", "Dépenses", "Net Entrées + Labo", "Net Pharmacie + Écho",
               "Recette nette", "Argent compté", "Écart", "Statut"]
    rows = []
    for s, st in d["days"]:
        rows.append([s.date, st.entries_excl_lab, st.lab, st.echo, st.pharmacy, st.gross, st.commissions_paid,
                     st.expenses_total, st.net_entree, st.net_pharma, st.net_total, s.counted_cash, s.cash_gap,
                     s.get_status_display()])
    last = _table(ws, 4, headers, rows,
                  ["TOTAL", t.entries_excl_lab, t.lab, t.echo, t.pharmacy, t.gross, t.commissions_paid,
                   t.expenses_total, t.net_entree, t.net_pharma, t.net_total, None, None, ""],
                  widths=[12] + [14] * 12 + [11])
    for r in range(5, last + 1):
        ws.cell(r, 1).number_format = "DD/MM/YYYY"
    r = last + 2
    for label, value in [("Argent en caisse au début du mois", d["cash_start"]),
                         ("+ Recette nette du mois", t.net_total),
                         ("− Versements bancaires du mois", d["deposits_total"]),
                         ("= Argent en caisse à la fin du mois (théorique)", d["cash_end"])]:
        ws.cell(r, 1, label).font = Font(bold=True)
        c = ws.cell(r, 6, value)
        c.number_format = NUM0
        r += 1

    ws = wb.create_sheet("Détail par acte")
    _title(ws, "RECETTES PAR ACTE (montants)", sub)
    acts = d["acts"]
    headers = ["Date"] + [a.name for a in acts] + ["Pharmacie", "Total"]
    rows = []
    for s, st in d["days"]:
        rows.append([s.date] + [st.acts_amount.get(a.id, 0) for a in acts] + [st.pharmacy, st.gross])
    last = _table(ws, 4, headers, rows,
                  ["TOTAL"] + [t.acts_amount.get(a.id, 0) for a in acts] + [t.pharmacy, t.gross],
                  widths=[12] + [12] * (len(acts) + 2))
    for r in range(5, last + 1):
        ws.cell(r, 1).number_format = "DD/MM/YYYY"
    r = last + 1
    ws.cell(r, 1, "Nombre / mois").font = Font(bold=True)
    for i, a in enumerate(acts, start=2):
        if a.by_quantity:
            ws.cell(r, i, t.acts_qty.get(a.id, 0)).number_format = NUM0

    ws = wb.create_sheet("Activité")
    _title(ws, "ACTIVITÉ DU MOIS", sub)
    _table(ws, 4, ["Acte", "Nombre", "Tarif actuel", "Montant"],
           [[x["act"].name, x["qty"] if x["act"].by_quantity else None, x["price"], x["amount"]]
            for x in d["activity"]],
           ["TOTAL", None, None, t.acts_total], widths=[36, 12, 14, 16], zero=True)
    r = ws.max_row + 2
    ws.cell(r, 1, "Commissions laboratoire").font = Font(bold=True, size=12)
    _table(ws, r + 1, ["Bénéficiaire", "Traitement", "", "Montant"],
           [[name, "Conservé en caisse" if c["retained"] else "Reversé", None, c["amount"]]
            for name, c in t.commissions.items()], widths=[36, 20, 4, 16])

    ws = wb.create_sheet("Dépenses")
    _title(ws, "DÉPENSES DU MOIS", sub)
    last = _table(ws, 4, ["Date", "Catégorie", "Caisse", "Motif", "Bénéficiaire", "Montant"],
                  [[x.session.date, x.category.name, x.get_fund_display(), x.description, x.beneficiary, x.amount]
                   for x in d["expenses"]],
                  ["TOTAL", "", "", "", "", t.expenses_total], widths=[12, 34, 24, 50, 20, 14], num_from=5)
    for r in range(5, last + 1):
        ws.cell(r, 1).number_format = "DD/MM/YYYY"

    ws = wb.create_sheet("Versements")
    _title(ws, "VERSEMENTS BANCAIRES", sub)
    last = _table(ws, 4, ["Date", "Banque", "Bordereau", "Déposant", "Montant"],
                  [[x.date, str(x.account), x.reference, x.depositor, x.amount] for x in d["deposits"]],
                  ["TOTAL", "", "", "", d["deposits_total"]], widths=[12, 28, 18, 20, 14], num_from=4)
    for r in range(5, last + 1):
        ws.cell(r, 1).number_format = "DD/MM/YYYY"
    return _save(wb)


def annual_workbook(d):
    wb = Workbook()
    sub = f"{d['facility'].name} – Année {d['year']}"
    g = d["grand"]
    ws = wb.active
    ws.title = "Par mois"
    _title(ws, "RÉCAPITULATIF ANNUEL", sub)
    headers = ["Mois", "Jours saisis", "Entrées (hors labo)", "Laboratoire", "Échographie", "Pharmacie",
               "Recette brute", "Commissions reversées", "Dépenses", "Recette nette", "Versements"]
    rows = [[r["name"], r["t"].sessions, r["t"].entries_excl_lab, r["t"].lab, r["t"].echo, r["t"].pharmacy,
             r["t"].gross, r["t"].commissions_paid, r["t"].expenses_total, r["t"].net_total, r["deposits"]]
            for r in d["rows"]]
    _table(ws, 4, headers, rows,
           ["TOTAL", g.sessions, g.entries_excl_lab, g.lab, g.echo, g.pharmacy, g.gross, g.commissions_paid,
            g.expenses_total, g.net_total, d["deposits_total"]], widths=[14] + [14] * 10)

    ws = wb.create_sheet("Activité")
    _title(ws, "NOMBRE D'ACTES PAR MOIS", sub)
    _table(ws, 4, ["Acte"] + d["month_names"] + ["Total", "Montant"],
           [[x["act"].name] + (x["months"] if x["act"].by_quantity else [None] * 12)
            + [x["total"] if x["act"].by_quantity else None, x["amount"]] for x in d["activity"]],
           widths=[34] + [8] * 12 + [10, 14])

    ws = wb.create_sheet("Dépenses")
    _title(ws, "DÉPENSES PAR CATÉGORIE ET PAR MOIS", sub)
    _table(ws, 4, ["Catégorie"] + d["month_names"] + ["Total"],
           [[cat] + vals + [tot] for cat, vals, tot in d["expenses"]],
           ["TOTAL"] + [sum(r[1][i] for r in d["expenses"]) for i in range(12)] + [g.expenses_total],
           widths=[40] + [10] * 12 + [12])
    return _save(wb)
