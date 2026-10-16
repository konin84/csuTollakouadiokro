"""Importe l'ancien « TABLEAU DE BORD RECETTES FSU » (une feuille par mois).

    python manage.py import_excel chemin\\vers\\TABLEAU_DE_BORD_RECETTES_FSU_2026.xlsx --annee 2026
    python manage.py import_excel fichier.xlsx --annee 2026 --essai      (n'enregistre rien)
    python manage.py import_excel fichier.xlsx --annee 2026 --remplacer  (supprime d'abord les journées importées)

Ce que fait l'import :
- lit chaque feuille mensuelle (repérée par la date en B4 et l'en-tête « Date » en A8) ;
- reprend la date de chaque ligne avec le mois de la feuille (corrige les dates 2025 de février, etc.) ;
- convertit les montants par acte en nombre d'actes (montant ÷ prix unitaire) ; si le montant ne
  tombe pas juste, il est conservé tel quel et signalé ;
- crée les dépenses (ristournes, divers pharmacie, dépenses 1 à 3) avec le motif tiré des commentaires ;
- crée les versements bancaires (dates corrigées et signalées) ;
- compare, mois par mois, la recette nette recalculée avec la colonne « TOTAL DES RECETTE » (AJ) d'Excel.
"""
import re
import unicodedata
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from caisse.models import (FUND_ENTREE, FUND_PHARMA, Act, ActEntry, ActPrice, BankAccount, CashSession, Deposit,
                           Expense, ExpenseCategory, PharmacySale)
from caisse.services import clear_caches, legacy_lab_act, month_bounds, period_totals
from core.audit import log
from core.models import Facility


def norm(text) -> str:
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", text).strip().lower()


HEADER_TO_ACT = {
    "consultations curatives": "consultation-curative",
    "consultations eleves": "consultation-eleves",
    "cpn": "cpn",
    "accouchements": "accouchement",
    "accouchements a domicile": "accouchement-domicile",
    "mise + pause": "mise-pose",
    "consultations gynecologiques": "consultation-gyneco",
    "echographie": "echographie",
    "carnet mere enfant": "carnet-mere-enfant",
    "carnet individuel": "carnet-individuel",
    "pansement": "pansement",
    "controle ta": "controle-ta",
    "injections": "injections",
    "gant": "gant",
    "petite chirurgie": "petite-chirurgie",
    "bandelette urinaire": "bandelette-urinaire",
    "ambulance": "ambulance",
    "avortement spontane": "avortement-spontane",
    "laboratoire recette": "laboratoire",
}

# en-tête -> (catégorie de dépense par défaut, caisse)
HEADER_TO_EXPENSE = {
    "depense ristne medecin": ("Ristourne médecin (échographie)", FUND_PHARMA),
    "depense ristne sfde": ("Ristourne SFDE (échographie)", FUND_PHARMA),
    "depenses pharmacie divers": ("Dépenses pharmacie divers", FUND_PHARMA),
    "depense 1": (None, FUND_ENTREE),
    "depense 2": (None, FUND_ENTREE),
    "depense 3": (None, FUND_ENTREE),
}

KEYWORDS = [
    (("internet", "reabonnement", "credit", "telephone"), "Internet et téléphone"),
    (("espace vert", "entretien", "climatiseur", "climatiseurs", "poubelle", "nettoyage", "ramassage"),
     "Entretien (espaces verts, nettoyage…)"),
    (("rame", "encre", "toner", "registre", "piles", "stylo", "bic"), "Fournitures de bureau (rames, encre…)"),
    (("carnet", "recu"), "Carnets et reçus"),
    (("carburant", "craburant", "essence", "gasoil", "transport", "immatriculation"), "Carburant et transport"),
    (("alcool", "coton", "solute", "perfuseur", "sparadrap", "compresse", "gant", "suture", "gel", "gobelet",
      "film echo"), "Consommables médicaux (alcool, coton, solutés…)"),
    (("materiel", "tensiometre", "batterie", "dijonteur", "disjoncteur"), "Matériel médical"),
    (("produit pharmaceutique", "produits pharmaceutiques", "medicament"), "Dépenses pharmacie divers"),
]


def guess_category(text: str) -> str:
    t = norm(text)
    for words, cat in KEYWORDS:
        if any(w in t for w in words):
            return cat
    return "Autres dépenses"


def describe_expense(comment: str, amount: int) -> str:
    """Retrouve dans le commentaire le libellé associé à un montant (« achat rame :9000 »)."""
    if not comment:
        return ""
    digits = str(amount)
    pattern = r"\s?".join(digits)  # tolère « 72 621 »
    m = re.search(r"(?<!\d)" + pattern + r"(?!\d)", comment)
    if not m:
        return ""
    before = comment[: m.start()]
    # coupe au montant précédent
    prev = list(re.finditer(r"\d[\d ]{2,}\d(?!\d)", before))
    if prev:
        before = before[prev[-1].end():]
    before = re.sub(r"^[\s\d\-–é.)]+", "", before)
    return before.strip(" :;-–,").strip()


def as_int(value) -> int:
    if value in (None, "", " "):
        return 0
    try:
        return int(round(float(value)))
    except (TypeError, ValueError):
        return 0


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        parts = [p for p in re.split(r"\D+", value) if p]
        if len(parts) == 3:
            d, m, y = (int(p) for p in parts)
            if y < 100:
                y += 2000
            try:
                return date(y, m, d)
            except ValueError:
                return None
    return None


class Command(BaseCommand):
    help = "Importe le tableau de bord Excel (feuilles mensuelles) dans l'application."

    def add_arguments(self, parser):
        parser.add_argument("fichier")
        parser.add_argument("--annee", type=int, required=True, help="Année à importer (ex. 2026)")
        parser.add_argument("--essai", action="store_true", help="Simulation : n'enregistre rien")
        parser.add_argument("--remplacer", action="store_true",
                            help="Supprime d'abord les journées et versements déjà importés pour cette année")
        parser.add_argument("--produits", action="store_true",
                            help="Importe aussi la liste de médicaments de la feuille « 1ER COMMANDE MEDICAMENT »")

    def handle(self, *args, **opts):
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover
            raise CommandError("openpyxl manquant") from exc
        path = Path(opts["fichier"])
        if not path.exists():
            raise CommandError(f"Fichier introuvable : {path}")
        self.year = opts["annee"]
        self.notes = []
        wb_values = openpyxl.load_workbook(path, data_only=True)

        facility = Facility.current()
        months = self.find_month_sheets(wb_values)
        if not months:
            raise CommandError("Aucune feuille mensuelle trouvée pour cette année.")

        try:
            with transaction.atomic():
                if opts["remplacer"]:
                    start, end = date(self.year, 1, 1), date(self.year, 12, 31)
                    n = CashSession.objects.filter(facility=facility, imported=True, date__range=(start, end))
                    Deposit.objects.filter(facility=facility, date__range=(start, date(self.year + 1, 1, 31)),
                                           note__startswith="Import Excel").delete()
                    n.delete()
                summary = []
                for month, ws in months:
                    summary.append(self.import_month(facility, month, ws))
                if opts["produits"]:
                    self.import_products(wb_values)
                if facility.opening_date is None or facility.opening_date > date(self.year, months[0][0], 1):
                    facility.opening_date = date(self.year, months[0][0], 1)
                    facility.save()
                self.report(summary)
                log("import", None, {"fichier": path.name, "annee": self.year, "mois": [m for m, _ in months]},
                    model="import Excel", object_repr=f"Import {path.name} ({self.year})")
                if opts["essai"]:
                    raise _DryRun()
        except _DryRun:
            self.stdout.write(self.style.WARNING("Mode essai : aucune donnée enregistrée."))

    # ------------------------------------------------------------------
    def find_month_sheets(self, wb):
        found = {}
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            b4 = ws["B4"].value
            if isinstance(b4, datetime) and b4.year == self.year and norm(ws["A8"].value) == "date":
                found.setdefault(b4.month, ws)
        return sorted(found.items())

    def header_map(self, ws):
        cols = {}
        for c in range(1, ws.max_column + 1):
            h = norm(ws.cell(8, c).value)
            if h:
                cols.setdefault(h, c)
        return cols

    def find_row(self, ws, label, col=1, start=30, end=70):
        for r in range(start, min(end, ws.max_row) + 1):
            if norm(ws.cell(r, col).value).startswith(label):
                return r
        return None

    def import_month(self, facility, month, ws):
        cols = self.header_map(ws)
        price_row = self.find_row(ws, "prix unitaire")
        total_row = self.find_row(ws, "montant total")
        start, end = month_bounds(self.year, month)
        label = f"{start:%m/%Y} ({ws.title.strip()})"
        acts = {a.code: a for a in Act.objects.all()}

        # 1) lignes journalières
        rows = []
        for r in range(9, (total_row or 41)):
            raw_date = ws.cell(r, 1).value
            if not isinstance(raw_date, (datetime, date)):
                continue
            try:
                day = date(self.year, month, raw_date.day)
            except ValueError:
                continue
            rows.append((r, day, raw_date))

        # 2) tarifs : prix de la feuille, contrôlés sur les montants réellement saisis
        prices = {}
        for header, code in HEADER_TO_ACT.items():
            c = cols.get(header)
            act = acts.get(code)
            if not c or not act or not act.by_quantity:
                continue
            values = [as_int(ws.cell(r, c).value) for r, _, _ in rows]
            values = [v for v in values if v]
            sheet_price = as_int(ws.cell(price_row, c).value) if price_row else 0
            current = act.price_on(start) or 0
            if not values:  # aucun acte ce mois-ci : on ne change pas le tarif
                prices[code] = current or sheet_price
                continue
            candidates = [p for p in (sheet_price, current) if p]
            chosen = None
            for p in candidates:
                if all(v % p == 0 for v in values):
                    chosen = p
                    break
            if chosen is None:  # aucun prix ne tombe juste : on garde le tarif en vigueur
                chosen = current or sheet_price
            if sheet_price and chosen != sheet_price and values:
                self.notes.append(f"{label} : prix « {act.name} » du tableau ({sheet_price} F) incohérent avec "
                                  f"les montants saisis → {chosen} F retenu")
            if chosen and chosen != current:
                ActPrice.objects.update_or_create(act=act, valid_from=start, defaults={"price": chosen})
                self.notes.append(f"{label} : tarif « {act.name} » {current} F → {chosen} F")
            prices[code] = chosen

        # 3) comptes bancaires selon les en-têtes
        deposits = self.read_deposits(ws, month, label)

        comment_cols = [c for h, c in cols.items() if h.startswith("commentaire")]
        comment_cols = list(range(min(comment_cols), ws.max_column + 1)) if comment_cols else []
        categories = {c.name: c for c in ExpenseCategory.objects.all()}
        excel_net = 0
        imported_days = 0
        wrong_dates = []
        for r, day, raw_date in rows:
            values = {h: as_int(ws.cell(r, c).value) for h, c in cols.items()}
            has_data = any(values.get(h) for h in list(HEADER_TO_ACT) + list(HEADER_TO_EXPENSE) + ["pharmacie"])
            if not has_data:
                continue
            if raw_date.month != month or raw_date.year != self.year:
                wrong_dates.append(raw_date)
            excel_net += as_int(ws.cell(r, cols["total des recette"]).value) if "total des recette" in cols else 0
            comment = " ".join(str(ws.cell(r, c).value).strip() for c in comment_cols
                               if isinstance(ws.cell(r, c).value, str) and ws.cell(r, c).value.strip())
            session, created = CashSession.objects.get_or_create(
                facility=facility, date=day,
                defaults={"status": CashSession.VALIDATED, "imported": True, "notes": ""},
            )
            if not created:
                self.notes.append(f"{label} : journée du {day:%d/%m/%Y} déjà présente, ignorée")
                continue
            session._audit_skip = True
            imported_days += 1
            for header, code in HEADER_TO_ACT.items():
                amount = values.get(header, 0)
                if not amount:
                    continue
                act = acts.get(code)
                if act is None and code == "laboratoire":
                    act = acts[code] = legacy_lab_act()
                if not act.by_quantity:
                    e = ActEntry(session=session, act=act, quantity=0, amount=amount, manual_amount=True)
                else:
                    price = prices.get(code) or 0
                    if price and amount % price == 0:
                        e = ActEntry(session=session, act=act, quantity=amount // price, unit_price=price)
                    else:
                        e = ActEntry(session=session, act=act, quantity=round(amount / price) if price else 0,
                                     unit_price=price or None, amount=amount, manual_amount=True,
                                     note=f"Montant Excel {amount} F non divisible par {price} F")
                        self.notes.append(f"{label} : {day:%d/%m} « {act.name} » {amount} F ne correspond pas "
                                          f"au prix {price} F (montant conservé)")
                e._audit_skip = True
                e.save()
            pharma = values.get("pharmacie", 0)
            if pharma:
                s = PharmacySale(session=session, description="Ventes pharmacie (total du jour, import Excel)",
                                 quantity=1, unit_price=pharma)
                s._audit_skip = True
                s.save()
            for header, (cat_name, fund) in HEADER_TO_EXPENSE.items():
                amount = values.get(header, 0)
                if not amount:
                    continue
                desc = describe_expense(comment, amount)
                name = cat_name or guess_category(desc or comment)
                category = categories.get(name) or categories["Autres dépenses"]
                x = Expense(session=session, category=category, fund=fund, amount=amount,
                            description=(desc or comment or f"{header} (import Excel)")[:255])
                x._audit_skip = True
                x.save()
            if comment and not any(values.get(h) for h in HEADER_TO_EXPENSE):
                session.notes = comment
                session.save(update_fields=["notes"])

        if wrong_dates:
            self.notes.append(f"{label} : {len(wrong_dates)} lignes avaient une date erronée "
                              f"(ex. {wrong_dates[0]:%d/%m/%Y}) → date reprise du mois de la feuille")
        for d in deposits:
            Deposit.objects.create(facility=facility, from_till=True, **d)

        clear_caches()
        total, _days = period_totals(start, end, facility)
        return {
            "label": label,
            "days": imported_days,
            "app_net": total.net_total,
            "excel_net": excel_net,
            "gross": total.gross,
            "deposits": sum(d["amount"] for d in deposits),
        }

    def read_deposits(self, ws, month, label):
        out = []
        header_row = None
        for r in range(35, min(ws.max_row, 80) + 1):
            if norm(ws.cell(r, 8).value).startswith("banque"):
                header_row = r
                break
        if not header_row:
            return out
        second_bank = norm(ws["AI5"].value)
        second_name = "AFG Bank" if "afg" in second_bank else "Banque Populaire"
        blocks = [(9, 10, "BOA", ""), (14, 15, second_name, FUND_ENTREE)]
        for date_col, amount_col, bank, fund in blocks:
            account, _ = BankAccount.objects.get_or_create(bank=bank, name=bank, defaults={"fund": fund})
            raw = []
            for r in range(header_row + 1, header_row + 7):
                if not norm(ws.cell(r, date_col - 1).value).startswith("versement"):
                    continue
                amount = as_int(ws.cell(r, amount_col).value)
                if not amount:
                    continue
                raw.append((r, parse_date(ws.cell(r, date_col).value), amount))
            for i, (r, d, amount) in enumerate(raw):
                fixed = d
                if d is None:
                    fixed = month_bounds(self.year, month)[1]
                    self.notes.append(f"{label} : versement {bank} de {amount} F sans date lisible → "
                                      f"{fixed:%d/%m/%Y}")
                elif d.year != self.year and not (month == 12 and d.year == self.year + 1):
                    fixed = d.replace(year=self.year)
                    nxt = next((x[1] for x in raw[i + 1:] if x[1] and x[1].year == self.year), None)
                    if (nxt and fixed > nxt) or fixed > month_bounds(self.year, month)[1].replace(day=28) \
                            and fixed.month > month + 1:
                        fixed = fixed.replace(month=month)
                    self.notes.append(f"{label} : versement {bank} de {amount} F daté {d:%d/%m/%Y} → "
                                      f"corrigé en {fixed:%d/%m/%Y}")
                out.append({"date": fixed, "account": account, "amount": amount,
                            "reference": f"Versement {i + 1}", "note": f"Import Excel – {ws.title.strip()}"})
        return out

    def import_products(self, wb):
        from stock.models import Product

        ws = next((w for w in wb.worksheets if "commande" in norm(w.title)), None)
        if ws is None:
            self.notes.append("Feuille de commande de médicaments introuvable")
            return
        n = 0
        for r in range(3, ws.max_row + 1):
            name = ws.cell(r, 1).value
            if not name or norm(name) == "total":
                continue
            price = as_int(ws.cell(r, 3).value)
            _, created = Product.objects.get_or_create(name=str(name).strip().upper(),
                                                       defaults={"purchase_price": price})
            n += created
        self.notes.append(f"{n} produits créés depuis la feuille « {ws.title} » (stock à saisir par inventaire)")

    def report(self, summary):
        self.stdout.write("")
        self.stdout.write("Mois           Jours   Recette nette app   Total Excel (AJ)        Écart     Versements")
        for s in summary:
            gap = s["app_net"] - s["excel_net"]
            line = (f"{s['label'][:14]:14} {s['days']:5}   {s['app_net']:>17,}   {s['excel_net']:>16,}"
                    f"   {gap:>10,}   {s['deposits']:>12,}").replace(",", " ")
            self.stdout.write(self.style.WARNING(line) if gap else line)
        if self.notes:
            self.stdout.write("\nCorrections et points à vérifier :")
            for n in self.notes:
                self.stdout.write(f"  - {n}")


class _DryRun(Exception):
    pass
